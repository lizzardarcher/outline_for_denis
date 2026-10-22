
"""Когорта прямых рефералов: регистрации и оплаты за период."""
from datetime import datetime, timedelta
from decimal import Decimal

from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook

from bot.models import ReferralTransaction, TelegramReferral, TelegramUser, Transaction

SUCCESS_Q = Q(paid=True) | Q(status="succeeded")
WITHDRAWAL_SIDE = "Вывод средств"
_BIGINT_MAX = 9223372036854775807


def search_referrers(query, limit=25):
    query = (query or "").strip()
    if not query:
        return []
    filters = (
        Q(username__icontains=query)
        | Q(first_name__icontains=query)
        | Q(last_name__icontains=query)
    )
    if query.isdigit():
        value = int(query)
        if value <= _BIGINT_MAX:
            filters |= Q(user_id=value)
    return list(
        TelegramUser.objects.filter(filters)
        .annotate(
            direct_referrals=Count(
                "given_referrals",
                filter=Q(given_referrals__level=1),
            )
        )
        .order_by("-direct_referrals", "username", "user_id")[:limit]
    )


def parse_period(params, today=None):
    today = today or timezone.localdate()
    preset = (params.get("preset") or "").strip()
    if preset == "7":
        return today - timedelta(days=6), today, None
    if preset == "30":
        return today - timedelta(days=29), today, None
    if preset == "month":
        return today.replace(day=1), today, None

    raw_from = (params.get("date_from") or "").strip()
    raw_to = (params.get("date_to") or "").strip()
    if not raw_from and not raw_to:
        return today - timedelta(days=29), today, None
    try:
        date_from = datetime.strptime(raw_from, "%Y-%m-%d").date() if raw_from else today - timedelta(days=29)
        date_to = datetime.strptime(raw_to, "%Y-%m-%d").date() if raw_to else today
    except ValueError:
        return None, None, "Укажите даты в формате ГГГГ-ММ-ДД."
    if date_from > date_to:
        return None, None, "Дата начала позже даты конца."
    return date_from, date_to, None


def _as_local(ts):
    if ts is None:
        return None
    if timezone.is_aware(ts):
        return timezone.localtime(ts)
    return ts


def build_report(referrer, date_from, date_to):
    """
    Регистрации — прямые рефералы с join_date в периоде.
    Оплаты и реферальные начисления — только с датой в выбранном периоде
    (по всем прямым рефералам, не только зарегистрированным в периоде).
    """
    all_direct = list(
        TelegramReferral.objects.filter(referrer=referrer, level=1)
        .select_related("referred")
        .order_by("referred__join_date", "referred_id")
    )
    cohort_users = [
        row.referred for row in all_direct
        if row.referred.join_date and date_from <= row.referred.join_date <= date_to
    ]
    cohort_ids = {user.pk for user in cohort_users}
    all_referred_ids = [row.referred_id for row in all_direct]
    referred_by_id = {row.referred_id: row.referred for row in all_direct}

    payments = []
    if all_referred_ids:
        for tx in (
            Transaction.objects.filter(SUCCESS_Q, user_id__in=all_referred_ids)
            .exclude(side=WITHDRAWAL_SIDE)
            .order_by("timestamp", "id")
        ):
            day = _as_local(tx.timestamp).date()
            if date_from <= day <= date_to:
                payments.append(tx)

    first_payment = {}
    payments_by_day = {}
    payment_count = 0
    payment_sum = Decimal("0")
    payer_ids_in_scope = set()
    for tx in payments:
        payment_count += 1
        amount = tx.amount or Decimal("0")
        payment_sum += amount
        day = _as_local(tx.timestamp).date()
        bucket = payments_by_day.setdefault(day, {"count": 0, "sum": Decimal("0")})
        bucket["count"] += 1
        bucket["sum"] += amount
        payer_ids_in_scope.add(tx.user_id)
        if tx.user_id not in first_payment:
            first_payment[tx.user_id] = tx

    referral_income_count = 0
    referral_income_sum = Decimal("0")
    if all_referred_ids:
        for rt in (
            ReferralTransaction.objects.filter(
                referral__referrer=referrer,
                referral__level=1,
            )
            .select_related("transaction", "referral")
            .order_by("timestamp", "id")
        ):
            day = _as_local(rt.timestamp).date()
            tx_day = (
                _as_local(rt.transaction.timestamp).date()
                if rt.transaction_id and rt.transaction
                else None
            )
            if date_from <= day <= date_to or (tx_day and date_from <= tx_day <= date_to):
                referral_income_count += 1
                referral_income_sum += rt.amount or Decimal("0")

    # Строки: зарегистрированные в периоде + прямые рефералы с оплатой в периоде.
    row_users = list(cohort_users)
    seen = set(cohort_ids)
    for user_id in sorted(payer_ids_in_scope):
        if user_id in seen:
            continue
        user = referred_by_id.get(user_id)
        if user is None:
            continue
        row_users.append(user)
        seen.add(user_id)

    regs_by_day = {}
    paid_by_join = {}
    rows = []
    for user in row_users:
        in_cohort = user.pk in cohort_ids
        if in_cohort:
            regs_by_day[user.join_date] = regs_by_day.get(user.join_date, 0) + 1
        first = first_payment.get(user.pk)
        if in_cohort and first:
            paid_by_join[user.join_date] = paid_by_join.get(user.join_date, 0) + 1
        rows.append({
            "user": user,
            "paid": first is not None,
            "first_at": _as_local(first.timestamp) if first else None,
            "first_amount": first.amount if first else None,
        })

    registrations = len(cohort_users)
    paid_users = sum(1 for user in cohort_users if user.pk in first_payment)
    if registrations:
        conversion = (Decimal(paid_users) * Decimal("100") / Decimal(registrations)).quantize(Decimal("0.1"))
    else:
        conversion = Decimal("0.0")

    daily = []
    cursor = date_from
    while cursor <= date_to:
        pay = payments_by_day.get(cursor, {"count": 0, "sum": Decimal("0")})
        daily.append({
            "date": cursor,
            "registrations": regs_by_day.get(cursor, 0),
            "paid_users": paid_by_join.get(cursor, 0),
            "payment_count": pay["count"],
            "payment_sum": pay["sum"],
        })
        cursor += timedelta(days=1)

    return {
        "referrer": referrer,
        "date_from": date_from,
        "date_to": date_to,
        "registrations": registrations,
        "paid_users": paid_users,
        "conversion_pct": conversion,
        "payment_count": payment_count,
        "payment_sum": payment_sum,
        "referral_income_count": referral_income_count,
        "referral_income_sum": referral_income_sum,
        "rows": rows,
        "daily": daily,
        "chart_labels": [day["date"].strftime("%d.%m") for day in daily],
        "chart_registrations": [day["registrations"] for day in daily],
        "chart_paid": [day["paid_users"] for day in daily],
    }


def excel_response(report):
    referrer = report["referrer"]
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Реферер", referrer.first_name or "", referrer.last_name or "", referrer.username or "", referrer.user_id])
    summary.append(["Период", report["date_from"].isoformat(), report["date_to"].isoformat()])
    summary.append(["Регистрации", report["registrations"]])
    summary.append(["Оплатили", report["paid_users"]])
    summary.append(["Конверсия, %", float(report["conversion_pct"])])
    summary.append(["Оплат", report["payment_count"]])
    summary.append(["Сумма", float(report["payment_sum"])])
    summary.append(["Реферальных начислений", report["referral_income_count"]])
    summary.append(["Сумма начислений", float(report["referral_income_sum"])])

    daily_ws = wb.create_sheet("По дням")
    daily_ws.append(["Дата", "Регистрации", "Из них оплатили", "Оплат за день", "Сумма за день"])
    for day in report["daily"]:
        daily_ws.append([
            day["date"].isoformat(),
            day["registrations"],
            day["paid_users"],
            day["payment_count"],
            float(day["payment_sum"]),
        ])

    people = wb.create_sheet("Люди")
    people.append([
        "user_id", "Имя", "Фамилия", "username", "Дата регистрации",
        "Оплатил", "Первая оплата в периоде", "Сумма",
    ])
    for row in report["rows"]:
        user = row["user"]
        first_at = row["first_at"]
        people.append([
            user.user_id,
            user.first_name or "",
            user.last_name or "",
            user.username or "",
            user.join_date.isoformat() if user.join_date else "",
            "да" if row["paid"] else "нет",
            first_at.strftime("%Y-%m-%d %H:%M") if first_at else "",
            float(row["first_amount"]) if row["first_amount"] is not None else "",
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"referral_{referrer.user_id}_{report['date_from'].isoformat()}_{report['date_to'].isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
