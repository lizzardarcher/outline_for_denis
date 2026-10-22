
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from bot.models import (
    IncomeInfo,
    ReferralTransaction,
    TelegramReferral,
    TelegramUser,
    Transaction,
)
from bot.referral_analytics import build_report


class ReferralCohortAnalyticsTest(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.referrer = TelegramUser.objects.create(
            user_id=1001, username="sharer", first_name="Sharer",
        )
        self.paid = TelegramUser.objects.create(
            user_id=1002, username="paiduser", first_name="Paid",
        )
        self.free = TelegramUser.objects.create(
            user_id=1003, username="freeuser", first_name="Free",
        )
        self.deep = TelegramUser.objects.create(
            user_id=1004, username="level2", first_name="Deep",
        )
        self.old = TelegramUser.objects.create(
            user_id=1005, username="olduser", first_name="Old",
        )
        for user in (self.paid, self.free, self.deep):
            TelegramUser.objects.filter(pk=user.pk).update(join_date=self.today)
        TelegramUser.objects.filter(pk=self.old.pk).update(
            join_date=self.today - timedelta(days=40),
        )
        self.ref_paid = TelegramReferral.objects.create(
            referrer=self.referrer, referred=self.paid, level=1,
        )
        TelegramReferral.objects.create(referrer=self.referrer, referred=self.free, level=1)
        TelegramReferral.objects.create(referrer=self.referrer, referred=self.deep, level=2)
        self.ref_old = TelegramReferral.objects.create(
            referrer=self.referrer, referred=self.old, level=1,
        )

        income = IncomeInfo.objects.create(id=1, total_amount=Decimal("0"))
        self.tx_paid = Transaction.objects.create(
            user=self.paid,
            amount=Decimal("500.00"),
            currency="RUB",
            side="Приход средств",
            status="succeeded",
            paid=True,
            income_info=income,
            description="Подписка",
        )
        Transaction.objects.create(
            user=self.paid,
            amount=Decimal("100.00"),
            currency="RUB",
            side="Вывод средств",
            status="succeeded",
            paid=True,
            income_info=income,
            description="Вывод",
        )
        Transaction.objects.create(
            user=self.deep,
            amount=Decimal("500.00"),
            currency="RUB",
            side="Приход средств",
            status="succeeded",
            paid=True,
            income_info=income,
        )
        self.tx_old = Transaction.objects.create(
            user=self.old,
            amount=Decimal("300.00"),
            currency="RUB",
            side="Приход средств",
            status="succeeded",
            paid=True,
            income_info=income,
            description="Подписка старого реферала",
        )
        ReferralTransaction.objects.create(
            referral=self.ref_paid,
            transaction=self.tx_paid,
            amount=Decimal("75.00"),
        )
        ReferralTransaction.objects.create(
            referral=self.ref_old,
            transaction=self.tx_old,
            amount=Decimal("45.00"),
        )

    def test_level_1_cohort_counts_payers_not_withdrawals(self):
        report = build_report(self.referrer, self.today, self.today)
        self.assertEqual(report["registrations"], 2)
        self.assertEqual(report["paid_users"], 1)

        # общие: когорта (500) + старый реферал (300)
        self.assertEqual(report["payment_count"], 2)
        self.assertEqual(report["payment_sum"], Decimal("800.00"))
        self.assertEqual(report["payment_count_new"], 1)
        self.assertEqual(report["payment_sum_new"], Decimal("500.00"))
        self.assertEqual(report["referral_income_count"], 2)
        self.assertEqual(report["referral_income_sum"], Decimal("120.00"))

        self.assertEqual(report["referral_income_count_new"], 1)
        self.assertEqual(report["referral_income_sum_new"], Decimal("75.00"))

        names_all = {row["user"].username for row in report["rows_all"]}
        self.assertEqual(names_all, {"paiduser", "olduser"})
        names_new = {row["user"].username for row in report["rows_new"]}
        self.assertEqual(names_new, {"paiduser"})

    def test_payment_outside_period_not_counted(self):
        yesterday = self.today - timedelta(days=1)
        TelegramUser.objects.filter(pk=self.paid.pk).update(join_date=yesterday)
        report = build_report(self.referrer, yesterday, yesterday)
        self.assertEqual(report["registrations"], 1)
        self.assertEqual(report["paid_users"], 0)
        self.assertEqual(report["payment_count"], 0)
        self.assertEqual(report["payment_sum"], Decimal("0"))
        self.assertEqual(report["rows_all"], [])
        self.assertEqual(report["rows_new"], [])

    def test_old_timestamp_excluded_from_period(self):
        """Платёж с timestamp вне периода не попадает в таблицу, даже если реферал активен."""
        old_ts = timezone.now() - timedelta(days=10)
        Transaction.objects.filter(pk=self.tx_paid.pk).update(timestamp=old_ts)
        Transaction.objects.filter(pk=self.tx_old.pk).update(timestamp=old_ts)
        ReferralTransaction.objects.all().update(timestamp=old_ts)
        report = build_report(self.referrer, self.today, self.today)
        self.assertEqual(report["payment_count"], 0)
        self.assertEqual(report["payment_sum"], Decimal("0"))
        self.assertEqual(report["rows_all"], [])
        self.assertEqual(report["referral_income_count"], 0)

    def test_older_referral_payment_in_period_counts(self):
        """Оплаты старых рефералов — только во вкладке общих платежей."""
        report = build_report(self.referrer, self.today, self.today)
        old_rows = [row for row in report["rows_all"] if row["user"].username == "olduser"]
        self.assertEqual(len(old_rows), 1)
        self.assertEqual(old_rows[0]["amount"], Decimal("300.00"))
        self.assertFalse(any(row["user"].username == "olduser" for row in report["rows_new"]))

    def test_admin_page_and_excel(self):
        User = get_user_model()
        admin = User.objects.create_superuser("refadmin", "ref@example.com", "pass")
        self.client.force_login(admin)
        page = reverse("admin:bot_telegramreferral_analytics")
        params = {
            "referrer": self.referrer.pk,
            "date_from": self.today.isoformat(),
            "date_to": self.today.isoformat(),
        }
        response = self.client.get(page, params, secure=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Paid")
        self.assertContains(response, "Платежи новые за период")
        self.assertContains(response, "Платежи общие за период")
        self.assertContains(response, "Новых за этот период")
        self.assertContains(response, "Оплат всех рефералов за этот период")
        self.assertContains(response, "Начислений от новых за этот период")
        self.assertNotContains(response, "Deep")

        changelist = self.client.get(reverse("admin:bot_telegramreferral_changelist"), secure=True)
        self.assertContains(changelist, "Расширенный анализ")

        export = self.client.get(reverse("admin:bot_telegramreferral_analytics_export"), params, secure=True)
        self.assertEqual(export.status_code, 200)
        self.assertIn("spreadsheetml", export["Content-Type"])
        book = load_workbook(BytesIO(export.content))
        summary = book["Сводка"]
        self.assertEqual(summary["A3"].value, "Новых за период")
        self.assertEqual(summary["B3"].value, 2)
        self.assertEqual(summary["A6"].value, "Оплат всех рефералов за период")
        self.assertEqual(summary["B6"].value, 2)
        self.assertEqual(summary["A10"].value, "Начислений от новых за период")
        self.assertEqual(summary["B10"].value, 1)
        self.assertEqual(summary["B11"].value, 75.0)
        self.assertEqual(summary["A12"].value, "Начислений всех за период")
        self.assertEqual(summary["B12"].value, 2)
        self.assertEqual(summary["B13"].value, 120.0)
        new_ws = book["Платежи новые"]
        all_ws = book["Платежи общие"]
        new_ids = {new_ws.cell(row=i, column=1).value for i in range(2, new_ws.max_row + 1)}
        all_ids = {all_ws.cell(row=i, column=1).value for i in range(2, all_ws.max_row + 1)}
        self.assertEqual(new_ids, {1002})
        self.assertEqual(all_ids, {1002, 1005})

    def test_support_cannot_open_analytics(self):
        from django.conf import settings

        User = get_user_model()
        support = User.objects.create_user(
            settings.SUPPORT_ACCOUNT, "support@example.com", "pass", is_staff=True,
        )
        self.client.force_login(support)
        response = self.client.get(reverse("admin:bot_telegramreferral_analytics"), secure=True)
        self.assertEqual(response.status_code, 403)
